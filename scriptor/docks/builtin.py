"""
Eingebaute Docks. Gleichzeitig Referenz-Implementierungen für eigene Plugins.

  lesbarkeit   Flesch-Amstad + Wiener Sachtextformel je Kapitel (Werkzeug, CLI, API)
  docx_export  Manuskript als Word-Datei (.docx) — nur aktiv, wenn python-docx installiert ist
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import TYPE_CHECKING, Any, Dict, List

from scriptor.docks import Dock, DockRegistry

try:  # Modulebene: FastAPI muss die Annotationen der Routen auflösen können
    from fastapi import APIRouter, Request
except ImportError:  # pragma: no cover - CLI ohne FastAPI
    APIRouter = Request = None

if TYPE_CHECKING:
    from scriptor.agent import Agent

# ---------------------------------------------------------------------------
# Lesbarkeit
# ---------------------------------------------------------------------------

_VOWELS = re.compile(r"[aeiouyäöü]+", re.I)
_WORD = re.compile(r"[A-Za-zÄÖÜäöüß]+")


def _syllables(word: str) -> int:
    return max(1, len(_VOWELS.findall(word)))


def readability(text: str) -> Dict[str, float]:
    """Deutsche Lesbarkeitsindizes (Näherung über Vokalgruppen als Silben)."""
    sentences = [s for s in re.split(r"[.!?…]+", text) if _WORD.search(s)]
    words = _WORD.findall(text)
    if not sentences or not words:
        return {"words": 0}
    syl = [_syllables(w) for w in words]
    asl = len(words) / len(sentences)
    asw = sum(syl) / len(words)
    ms = 100 * sum(1 for s in syl if s >= 3) / len(words)
    iw = 100 * sum(1 for w in words if len(w) > 6) / len(words)
    es = 100 * sum(1 for s in syl if s == 1) / len(words)
    return {
        "words": len(words),
        "avg_sentence_len": round(asl, 1),
        "flesch_amstad": round(180 - asl - 58.5 * asw, 1),
        "wiener_sachtext": round(0.1935 * ms + 0.1672 * asl + 0.1297 * iw - 0.0327 * es - 0.875, 1),
    }


def interpret(flesch: float) -> str:
    if flesch >= 80:
        return "sehr leicht (Erstleser/Bilderbuch)"
    if flesch >= 65:
        return "leicht (Kinderbuch)"
    if flesch >= 50:
        return "mittel (YA/Unterhaltung)"
    if flesch >= 30:
        return "anspruchsvoll (Sachbuch/Literatur)"
    return "schwer (Fachtext)"


def book_readability(project: Any) -> List[Dict[str, Any]]:
    rows = []
    for ch in project.outline:
        latest = project.latest_chapter(int(ch["nr"]))
        if latest:
            r: Dict[str, Any] = dict(readability(latest[1]))
            r["nr"] = int(ch["nr"])
            r["level"] = interpret(r.get("flesch_amstad", 0))
            rows.append(r)
    return rows


class LesbarkeitDock(Dock):
    name = "lesbarkeit"
    version = "1.0.0"
    description = "Lesbarkeitsindizes (Flesch-Amstad, Wiener Sachtextformel) je Kapitel"

    def load(self, registry: DockRegistry) -> bool:
        def tool(agent: "Agent", slug: str) -> str:
            rows = book_readability(agent._book(slug))
            if not rows:
                return "Noch keine Kapitel mit Text."
            return "\n".join(
                f"Kap. {r['nr']}: Flesch {r['flesch_amstad']} ({r['level']}), "
                f"WSTF {r['wiener_sachtext']}, Ø Satz {r['avg_sentence_len']} W."
                for r in rows
            )

        registry.add_tool(
            "readability",
            "Lesbarkeit aller Kapitel eines Buchs (Flesch-Amstad, Wiener Sachtextformel). "
            "Nutze es, um Zielgruppen-Passung zu prüfen.",
            {"type": "object", "properties": {"slug": {"type": "string"}},
             "required": ["slug"], "additionalProperties": False},
            tool,
        )

        try:
            import typer

            cli = typer.Typer(help="Lesbarkeit")

            @cli.command("buch")
            def _cli(slug: str) -> None:
                from scriptor.project import BookProject
                from scriptor.settings import Settings

                p = BookProject.load(Settings.load().workspace_path(), slug)
                for r in book_readability(p):
                    typer.echo(f"Kap. {r['nr']:>3}: Flesch {r['flesch_amstad']:>6} "
                               f"WSTF {r['wiener_sachtext']:>5}  {r['level']}")

            registry.add_cli("lesbarkeit", cli)
        except ImportError:
            pass

        if APIRouter is not None:
            router = APIRouter()

            @router.get("/api/docks/lesbarkeit/{slug}")
            def _api(slug: str, request: Request) -> List[Dict[str, Any]]:
                return book_readability(request.app.state.project(slug))

            registry.add_router(router)
        return True


# ---------------------------------------------------------------------------
# DOCX-Export
# ---------------------------------------------------------------------------


def export_docx(project: Any) -> Path:
    import docx  # python-docx

    md = project.export_markdown().read_text(encoding="utf-8")
    doc = docx.Document()
    for line in md.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_page_break()
            doc.add_heading(line[3:], level=1)
        elif line.startswith("*") and line.endswith("*") and len(line) > 2:
            doc.add_paragraph(line.strip("*")).italic = True
        elif line.strip():
            doc.add_paragraph(line)
    out = project.root / "export" / f"{project.slug}.docx"
    doc.save(str(out))
    return Path(out)


class DocxExportDock(Dock):
    name = "docx_export"
    version = "1.0.0"
    description = "Manuskript als Word-Datei (.docx) exportieren"

    def load(self, registry: DockRegistry) -> bool:
        try:
            import docx  # noqa: F401
        except ImportError:
            return False  # optional: pip install python-docx

        def tool(agent: "Agent", slug: str) -> str:
            out = export_docx(agent._book(slug))
            return f"Exportiert: {out.relative_to(agent.workspace)}"

        registry.add_tool(
            "export_docx",
            "Exportiert das Manuskript als Word-Datei (.docx).",
            {"type": "object", "properties": {"slug": {"type": "string"}},
             "required": ["slug"], "additionalProperties": False},
            tool,
        )
        return True


BUILTIN_DOCKS: List[Any] = [LesbarkeitDock, DocxExportDock]
